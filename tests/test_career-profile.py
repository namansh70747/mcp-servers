"""Offline tests for P2 profile.json adoption in the career cluster:
blog-drafter, deckforge, github-profile.

These verify the NEW additive behavior — auto-filling author/name/headline/contact/links
from the suite's shared profile.json — AND that every tool still works with NO profile.json
present (signatures and old behavior preserved). Fully offline: no network, no credentials.
GitHub-backed tools are exercised only via build_readme with network calls stubbed out.

Run with the suite venv:
    VIRTUAL_ENV= .venv/bin/python tests/test_career-profile.py
"""
from __future__ import annotations

import asyncio
import importlib.util
import os
import tempfile
from pathlib import Path

# Isolate all runtime data in a throwaway dir BEFORE importing any server.
os.environ["MCP_NO_DOTENV"] = "1"
_TMP = tempfile.mkdtemp(prefix="career-profile-test-")
os.environ["MCP_DATA_DIR"] = _TMP

from fastmcp import Client  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]

SAMPLE_PROFILE = {
    "name": "Ada Lovelace",
    "email": "ada@example.com",
    "phone": "+1-555-0100",
    "location": "London, UK",
    "links": {"github": "https://github.com/ada", "website": "https://ada.dev", "twitter": ""},
    "headline": "Backend engineer & analytical-engine enthusiast",
    "summary": "I build reliable systems and write about them.",
    "skills": {"languages": ["Python", "Rust"], "tools": ["Docker"]},
}


def _load(name: str, fname: str):
    path = ROOT / "servers" / name / "server.py"
    spec = importlib.util.spec_from_file_location(fname, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


async def _tools(mcp):
    async with Client(mcp) as c:
        return {t.name for t in await c.list_tools()}


async def _call(mcp, tool, args):
    async with Client(mcp) as c:
        return (await c.call_tool(tool, args)).data


def _with_profile(mod, profile: dict | None):
    """Point a freshly-loaded module's profile loader at a controlled dict (or empty)."""
    mod._profile = lambda: (profile or {})


# ------------------------------------------------------------------ blog-drafter
def test_blog_drafter_profile():
    mod = _load("blog-drafter", "bd_profile_t")
    mcp = mod.mcp

    # With a profile present, new_draft / outline_to_draft auto-fill the author byline.
    _with_profile(mod, SAMPLE_PROFILE)
    d = asyncio.run(_call(mcp, "outline_to_draft", {"topic": "scaling", "kind": "blog"}))
    assert d["author"] == "Ada Lovelace — Backend engineer & analytical-engine enthusiast"
    did = d["id"]
    full = asyncio.run(_call(mcp, "get", {"draft_id": did}))
    assert full["author"] == d["author"]
    # frontmatter export carries the author line
    md = asyncio.run(_call(mcp, "export_md", {"draft_id": did, "with_frontmatter": True}))
    text = Path(md["path"]).read_text()
    assert "author: Ada Lovelace — Backend engineer" in text

    # explicit author overrides the profile default
    d2 = asyncio.run(_call(mcp, "new_draft", {"title": "Hi", "author": "Someone Else"}))
    f2 = asyncio.run(_call(mcp, "get", {"draft_id": d2["id"]}))
    assert f2["author"] == "Someone Else"

    # "-" sentinel suppresses the byline in frontmatter
    d3 = asyncio.run(_call(mcp, "new_draft", {"title": "No byline post here for sure", "author": "-"}))
    md3 = asyncio.run(_call(mcp, "export_md", {"draft_id": d3["id"], "with_frontmatter": True}))
    assert "author:" not in Path(md3["path"]).read_text()

    # With NO profile, behavior is preserved: no author byline, draft still created.
    _with_profile(mod, {})
    d4 = asyncio.run(_call(mcp, "new_draft", {"title": "Plain draft"}))
    f4 = asyncio.run(_call(mcp, "get", {"draft_id": d4["id"]}))
    assert (f4.get("author") or "") == ""
    md4 = asyncio.run(_call(mcp, "export_md", {"draft_id": d4["id"], "with_frontmatter": True}))
    assert "author:" not in Path(md4["path"]).read_text()


# ------------------------------------------------------------------ deckforge
def test_deckforge_profile():
    mod = _load("deckforge", "df_profile_t")
    mcp = mod.mcp

    # With a profile, a career template fills the title slide name/headline + Contact slide.
    _with_profile(mod, SAMPLE_PROFILE)
    c = asyncio.run(_call(mcp, "create_presentation", {"template": "career", "theme": "midnight"}))
    did = c["deck_id"]
    info = asyncio.run(_call(mcp, "deck_info", {"deck_id": did}))
    titles = [s["title"] for s in info["outline"]]
    assert "Ada Lovelace" in titles  # title slide used the profile name
    # save + confirm the rendered pptx text contains the contact line
    out = asyncio.run(_call(mcp, "save_presentation", {"deck_id": did, "filename": "_career.pptx"}))
    assert os.path.exists(out["path"])
    from pptx import Presentation
    text = "\n".join(
        sh.text_frame.text
        for slide in Presentation(out["path"]).slides
        for sh in slide.shapes
        if sh.has_text_frame
    )
    assert "ada@example.com" in text and "London, UK" in text
    assert "Backend engineer" in text  # headline subtitle

    # Explicit title wins over the profile name (old behavior preserved).
    c2 = asyncio.run(_call(mcp, "create_presentation", {"template": "career", "title": "My Deck"}))
    info2 = asyncio.run(_call(mcp, "deck_info", {"deck_id": c2["deck_id"]}))
    assert info2["outline"][0]["title"] == "My Deck"

    # With NO profile, the career template falls back to the old placeholder behavior.
    _with_profile(mod, {})
    c3 = asyncio.run(_call(mcp, "create_presentation", {"template": "career"}))
    info3 = asyncio.run(_call(mcp, "deck_info", {"deck_id": c3["deck_id"]}))
    assert info3["outline"][0]["title"] == "Career"  # template name title-cased
    assert info3["slides"] == len(mod.TEMPLATES["career"])  # same slide count as before


# ------------------------------------------------------------------ github-profile
def test_github_profile_readme():
    mod = _load("github-profile", "ghp_profile_t")
    mcp = mod.mcp

    # Stub all network so build_readme runs fully offline. suggest_pins/language_stats call _get.
    mod._get = lambda url, **params: {"__error__": "offline"}
    # _all_repos returns a dict (error) -> suggest_pins/language_stats return that dict gracefully.

    _with_readme_profile(mod, SAMPLE_PROFILE)
    r = asyncio.run(_call(mcp, "build_readme", {"username": "ada", "write": False,
                                                "include_stats": False, "include_badges": False}))
    md = r["markdown"]
    assert "Ada Lovelace" in md                       # name from profile
    assert "Backend engineer" in md                   # headline
    assert "I build reliable systems" in md           # summary/bio
    assert "Python · Rust · Docker" in md             # skills flattened
    assert "[github](https://github.com/ada)" in md   # links pulled in
    assert "[website](https://ada.dev)" in md
    assert "twitter" not in md                         # empty link omitted

    # With NO profile.json, build_readme falls back to the username and still produces a README.
    _with_readme_profile(mod, None)
    r2 = asyncio.run(_call(mcp, "build_readme", {"username": "ada", "write": False,
                                                 "include_stats": False, "include_badges": False}))
    assert "Hi, I'm ada" in r2["markdown"]

    # Input validation still rejects malformed usernames offline (no network).
    bad = asyncio.run(_call(mcp, "build_readme", {"username": "bad name", "write": False}))
    assert "__error__" in bad


def _with_readme_profile(mod, profile: dict | None):
    """github-profile reads profile.json inline via ROOT/profile.json — redirect ROOT at a temp
    dir holding (or omitting) the controlled profile.json so the read path itself is exercised."""
    tmp = Path(tempfile.mkdtemp(prefix="ghp-prof-"))
    if profile is not None:
        import json
        (tmp / "profile.json").write_text(json.dumps(profile))
    mod.ROOT = tmp


# ------------------------------------------------------------------ shape guards
def test_get_shape_unchanged():
    """The httpx->mcp_base.http swap must keep _get's return shapes identical."""
    mod = _load("github-profile", "ghp_shape_t")

    # network failure -> request_failed envelope (status None path)
    mod.http.request = lambda *a, **k: {"ok": False, "error": "boom", "status": None}
    r = mod._get("https://api.github.com/x")
    assert r["__error__"] == "request_failed" and r["detail"] == "boom"

    # 404 -> numeric __error__ + hint
    mod.http.request = lambda *a, **k: {"ok": False, "status": 404, "text": "Not Found"}
    r = mod._get("https://api.github.com/x")
    assert r["__error__"] == 404 and "Not found" in r["hint"]

    # 403 rate limit -> hint
    mod.http.request = lambda *a, **k: {"ok": False, "status": 403, "text": "API rate limit exceeded"}
    r = mod._get("https://api.github.com/x")
    assert r["__error__"] == 403 and "Rate limited" in r["hint"]

    # 200 JSON -> parsed payload returned as-is
    mod.http.request = lambda *a, **k: {"ok": True, "status": 200, "text": "{}", "json": {"login": "ada"}}
    r = mod._get("https://api.github.com/x")
    assert r == {"login": "ada"}

    # 200 but non-JSON -> invalid_json
    mod.http.request = lambda *a, **k: {"ok": True, "status": 200, "text": "<html>"}
    r = mod._get("https://api.github.com/x")
    assert r["__error__"] == "invalid_json"


if __name__ == "__main__":
    for fn in (test_blog_drafter_profile, test_deckforge_profile,
               test_github_profile_readme, test_get_shape_unchanged):
        fn()
        print(fn.__name__, "OK")
    print("ALL CAREER-PROFILE TESTS PASSED")
