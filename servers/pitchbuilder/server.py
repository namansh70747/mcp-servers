"""pitchbuilder — stores the Claude-written, company-tailored project idea and renders it into
your FIXED outreach template. Emits a one-pager (markdown -> PDF via LibreOffice), manages a
reusable template library, suggests outreach angles, and keeps per-company A/B variants.

The agent writes the pitch (using funding-radar context + your profile.json); this server keeps
it per company and fills the template so every email follows your exact format.
"""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

import sqlite3

from jinja2 import Template
from mcp_base import BaseStore, data_dir, db_path, make_server

mcp = make_server(
    "pitchbuilder",
    instructions=("Store a tailored project pitch per company and render it into your template. "
                  "save_pitch -> render_into_template -> (optional) make_onepager -> onepager_to_pdf. "
                  "suggest_angles for ideas; save_variant for A/B; save_template for a reusable library."),
)

SCHEMA = """
CREATE TABLE IF NOT EXISTS pitches(
  id INTEGER PRIMARY KEY, company TEXT, domain TEXT, angle TEXT DEFAULT '',
  pitch TEXT NOT NULL, tags TEXT DEFAULT '', created_at TEXT, updated_at TEXT,
  UNIQUE(company)
);
CREATE TABLE IF NOT EXISTS pitch_variants(
  id INTEGER PRIMARY KEY, company TEXT, variant TEXT, angle TEXT DEFAULT '',
  pitch TEXT NOT NULL, created_at TEXT,
  UNIQUE(company, variant)
);
CREATE TABLE IF NOT EXISTS templates(
  id INTEGER PRIMARY KEY, name TEXT UNIQUE, subject TEXT DEFAULT '', body TEXT NOT NULL,
  tags TEXT DEFAULT '', created_at TEXT, updated_at TEXT
);
"""
store = BaseStore(db_path("pitchbuilder"), schema=SCHEMA)
OUT = data_dir("pitchbuilder")  # one-pagers land here
MAX_MD_BYTES = 5 * 1024 * 1024  # 5 MB cap when converting a markdown one-pager
ROOT = Path(__file__).resolve().parents[2]
PROFILE = ROOT / "profile.json"

ANGLES = [
    ("recent_funding", "They just raised — scaling pain is imminent.",
     "Congrats on the {round} round — as you scale, {pain} usually becomes the bottleneck; here's a concrete fix I could ship."),
    ("hiring_spike", "They're hiring fast — build-vs-buy and onboarding friction.",
     "Saw you're hiring across {team} — I sketched a tool that would cut ramp time for those new hires."),
    ("new_product", "They shipped something new — integration / polish opportunity.",
     "Loved the launch of {product} — here's a small addition I think would lift activation."),
    ("open_role", "There's a role you'd fit — show, don't tell.",
     "Rather than just applying for the {role} role, I built a small piece of what I'd own day one."),
    ("manual_workflow", "An obviously manual process you could automate.",
     "Noticed {workflow} still looks manual — I prototyped an automation that handles the tedious part."),
    ("data_leverage", "They sit on data they're underusing.",
     "You're collecting {data}; here's a lightweight way to turn it into a feature users would notice."),
]


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


class _SafeFill(dict):
    """str.format_map helper: leave unknown {placeholders} untouched instead of raising."""

    def __missing__(self, key: str) -> str:  # noqa: D401
        return "{" + key + "}"


def _ensure_columns() -> None:
    """Additively migrate older pitches rows (tags, updated_at)."""
    have = {r["name"] for r in store.query("PRAGMA table_info(pitches)")}
    for col, decl in {"tags": "TEXT DEFAULT ''", "updated_at": "TEXT"}.items():
        if col not in have:
            store.execute(f"ALTER TABLE pitches ADD COLUMN {col} {decl}")


_ensure_columns()


def _profile() -> dict:
    if PROFILE.exists():
        try:
            return json.loads(PROFILE.read_text())
        except Exception:  # noqa: BLE001
            return {}
    return {}


def _funding_context(company: str) -> dict:
    """Read-only lookup of the latest funding-radar lead for a company (round/amount/sector/domain).

    Opens ~/.mcp-suite/funding-radar/store.db read-only; returns {} if it is missing, locked, the
    schema differs, or the company is unknown. Never writes. Used to prefill {round}/{amount}/etc.
    placeholders so suggest_angles starters come back specific instead of generic.
    """
    if not company:
        return {}
    fr_db = db_path("funding-radar")
    if not Path(fr_db).exists():
        return {}
    conn = None
    try:
        conn = sqlite3.connect(f"file:{fr_db}?mode=ro", uri=True, timeout=2.0)
        conn.row_factory = sqlite3.Row
        cur = conn.execute(
            "SELECT company, domain, round, amount, amount_usd, sector, url, headline, found_at "
            "FROM leads WHERE lower(company)=? ORDER BY found_at DESC LIMIT 1",
            (company.strip().lower(),),
        )
        row = cur.fetchone()
        if row is None:
            # fall back to a fuzzy match on a leading substring
            cur = conn.execute(
                "SELECT company, domain, round, amount, amount_usd, sector, url, headline, found_at "
                "FROM leads WHERE lower(company) LIKE ? ORDER BY found_at DESC LIMIT 1",
                (f"{company.strip().lower()}%",),
            )
            row = cur.fetchone()
        return dict(row) if row else {}
    except Exception:  # noqa: BLE001 — missing DB / locked / schema drift -> no context
        return {}
    finally:
        if conn is not None:
            conn.close()


@mcp.tool
def read_profile() -> dict:
    """Return profile.json facts (name, headline, projects, skills, links) to ground a pitch."""
    return _profile()


@mcp.tool
def save_pitch(company: str, pitch: str, domain: str = "", angle: str = "", tags: str = "") -> dict:
    """Save (or replace) the tailored project pitch for a company."""
    now = _now()
    pid = store.execute(
        "INSERT INTO pitches(company,domain,angle,pitch,tags,created_at,updated_at) VALUES(?,?,?,?,?,?,?) "
        "ON CONFLICT(company) DO UPDATE SET pitch=excluded.pitch,domain=excluded.domain,"
        "angle=excluded.angle,tags=excluded.tags,updated_at=excluded.updated_at",
        (company, domain, angle, pitch, tags, now, now),
    )
    return {"id": pid, "company": company}


@mcp.tool
def get_pitch(company: str) -> dict:
    """Get the stored pitch for a company."""
    return store.query_one("SELECT * FROM pitches WHERE company=?", (company,)) or {"error": "no pitch"}


@mcp.tool
def list_pitches(limit: int = 50, tag: str = "") -> list[dict]:
    """List saved pitches, optionally filtered by tag."""
    if tag:
        return store.query("SELECT id, company, angle, tags, created_at FROM pitches WHERE tags LIKE ? "
                           "ORDER BY created_at DESC LIMIT ?", (f"%{tag}%", limit))
    return store.query("SELECT id, company, angle, tags, created_at FROM pitches ORDER BY created_at DESC LIMIT ?",
                       (limit,))


@mcp.tool
def render_into_template(template: str, context: dict) -> dict:
    """Render your fixed outreach template (Jinja2 string) with a context dict.

    Convention: put `Subject: ...` on the first line of the template; it's split out so the
    result is {subject, body}. Common context keys: name, company, role, pitch, sender, link.
    """
    rendered = Template(template).render(**context)
    subject, body = "", rendered
    if rendered.lower().startswith("subject:"):
        first, _, rest = rendered.partition("\n")
        subject = first.split(":", 1)[1].strip()
        body = rest.lstrip("\n")
    return {"subject": subject, "body": body}


@mcp.tool
def make_onepager(company: str, sections: dict, filename: str = "", tagline: str = "") -> dict:
    """Write a markdown one-pager (problem/solution/why-us/ask) to attach. Returns its path.
    `sections` = {problem, solution, how_it_works, impact, ask, ...}. Adds a contact line from
    profile.json. Convert to PDF with onepager_to_pdf()."""
    prof = _profile()
    title = f"# A project idea for {company}\n\n"
    if tagline:
        title += f"_{tagline}_\n\n"
    body = "\n\n".join(f"## {k.replace('_', ' ').title()}\n\n{v}" for k, v in sections.items())
    contact_bits = [prof.get("name", ""), prof.get("email", "")]
    links = prof.get("links", {}) or {}
    contact_bits += [v for v in (links.get("github"), links.get("website")) if v]
    contact = " · ".join(b for b in contact_bits if b)
    footer = f"\n\n---\n\n{contact}\n" if contact else "\n"
    md = title + body + footer
    # Use only the basename of any caller-supplied filename so a one-pager can never be written
    # outside this server's output dir (prevents '../' / absolute-path traversal).
    if filename:
        safe_name = Path(filename).name
        if not safe_name or safe_name in (".", ".."):
            return {"error": f"invalid filename: {filename!r}"}
    else:
        safe_name = f"onepager_{company.lower().replace(' ', '_')}.md"
    path = OUT / safe_name
    path.write_text(md, encoding="utf-8")
    return {"path": str(path), "bytes": len(md), "company": company}


@mcp.tool
def onepager_to_pdf(md_path: str) -> dict:
    """Convert a markdown one-pager to PDF: render md -> .docx (python-docx) -> .pdf (LibreOffice).
    Returns the PDF path, or an actionable error if LibreOffice (soffice) is not installed."""
    import shutil
    import subprocess
    src = Path(md_path).expanduser()
    if not src.is_file():
        return {"ok": False, "error": f"no file at {src}"}
    if src.stat().st_size > MAX_MD_BYTES:
        return {"ok": False, "error": f"file too large ({src.stat().st_size} bytes > {MAX_MD_BYTES})"}
    try:
        from docx import Document
    except ImportError:
        return {"ok": False, "error": "python-docx not installed"}
    doc = Document()
    for line in src.read_text(encoding="utf-8").splitlines():
        s = line.rstrip()
        if s.startswith("# "):
            doc.add_heading(s[2:], level=0)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=1)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=2)
        elif s.strip() == "---":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(s)
    docx_path = OUT / (src.stem + ".docx")
    doc.save(str(docx_path))
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {"ok": False, "docx": str(docx_path),
                "error": "LibreOffice not installed (brew install --cask libreoffice). DOCX written."}
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(docx_path)],
                       check=True, capture_output=True, timeout=120)
        return {"ok": True, "pdf": str(OUT / (src.stem + ".pdf")), "docx": str(docx_path)}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "docx": str(docx_path), "error": str(e)}


@mcp.tool
def list_onepagers() -> list[dict]:
    """List generated one-pager / PDF files in the output dir."""
    out = []
    for p in sorted(OUT.glob("onepager_*")) + sorted(OUT.glob("*.pdf")):
        out.append({"path": str(p), "name": p.name, "bytes": p.stat().st_size})
    return out


@mcp.tool
def suggest_angles(company: str = "", context: str = "") -> dict:
    """Suggest outreach angles for a company. Rules-based, free, offline — returns a list of
    {angle, rationale, starter} the agent can pick from and fill with specifics.

    If `company` matches a funding-radar lead (read-only ~/.mcp-suite/funding-radar/store.db),
    its round/amount/sector/domain are returned under `funding` and each angle also gets a
    `prefilled_starter` with {round}/{amount}/{sector} substituted. Old behavior (no company, or
    no matching lead) is unchanged: starters keep their raw {placeholders}."""
    funding = _funding_context(company) if company else {}
    fill = {
        "round": funding.get("round") or "",
        "amount": funding.get("amount") or "",
        "sector": funding.get("sector") or "",
    }
    out = []
    for key, rationale, starter in ANGLES:
        item = {"angle": key, "rationale": rationale, "starter": starter}
        if funding:
            # Best-effort prefill of the known placeholders; unknown ones stay as {placeholder}.
            try:
                item["prefilled_starter"] = starter.format_map(_SafeFill(fill))
            except Exception:  # noqa: BLE001
                item["prefilled_starter"] = starter
        out.append(item)
    result = {"company": company, "context": context, "angles": out,
              "_hint": "pick one, fill {placeholders}, then save_pitch with angle=<key>."}
    if funding:
        result["funding"] = {
            "company": funding.get("company") or company,
            "domain": funding.get("domain") or "",
            "round": funding.get("round") or "",
            "amount": funding.get("amount") or "",
            "amount_usd": funding.get("amount_usd"),
            "sector": funding.get("sector") or "",
            "headline": funding.get("headline") or "",
            "url": funding.get("url") or "",
            "source": "funding-radar",
        }
    return result


@mcp.tool
def save_variant(company: str, variant: str, pitch: str, angle: str = "") -> dict:
    """Save a per-company A/B pitch variant (e.g. 'short', 'technical'). Keyed by company+variant."""
    vid = store.execute(
        "INSERT INTO pitch_variants(company,variant,angle,pitch,created_at) VALUES(?,?,?,?,?) "
        "ON CONFLICT(company,variant) DO UPDATE SET pitch=excluded.pitch,angle=excluded.angle,"
        "created_at=excluded.created_at",
        (company, variant, angle, pitch, _now()))
    return {"id": vid, "company": company, "variant": variant}


@mcp.tool
def get_variant(company: str, variant: str) -> dict:
    """Get a specific per-company pitch variant."""
    return (store.query_one("SELECT * FROM pitch_variants WHERE company=? AND variant=?",
                            (company, variant)) or {"error": "no variant"})


@mcp.tool
def list_variants(company: str) -> list[dict]:
    """List all stored pitch variants for a company."""
    return store.query("SELECT id, variant, angle, created_at FROM pitch_variants WHERE company=? "
                       "ORDER BY variant", (company,))


@mcp.tool
def save_template(name: str, body: str, subject: str = "", tags: str = "") -> dict:
    """Store or replace a reusable email/pitch template in the managed library (keyed by name)."""
    now = _now()
    tid = store.execute(
        "INSERT INTO templates(name,subject,body,tags,created_at,updated_at) VALUES(?,?,?,?,?,?) "
        "ON CONFLICT(name) DO UPDATE SET subject=excluded.subject,body=excluded.body,"
        "tags=excluded.tags,updated_at=excluded.updated_at",
        (name, subject, body, tags, now, now))
    return {"id": tid, "name": name}


@mcp.tool
def get_template(name: str) -> dict:
    """Get a template from the managed library."""
    return store.query_one("SELECT * FROM templates WHERE name=?", (name,)) or {"error": "no template"}


@mcp.tool
def list_templates(tag: str = "") -> list[dict]:
    """List templates in the managed library, optionally filtered by tag."""
    if tag:
        return store.query("SELECT id, name, subject, tags FROM templates WHERE tags LIKE ? ORDER BY name",
                           (f"%{tag}%",))
    return store.query("SELECT id, name, subject, tags FROM templates ORDER BY name")


@mcp.tool
def delete_template(name: str) -> dict:
    """Delete a template from the managed library."""
    existed = store.query_one("SELECT 1 FROM templates WHERE name=?", (name,))
    store.execute("DELETE FROM templates WHERE name=?", (name,))
    return {"ok": True, "name": name, "deleted": bool(existed)}


@mcp.tool
def render_template(name: str, context: dict) -> dict:
    """Render a library template (Jinja2) with a context dict. Returns {subject, body}."""
    row = store.query_one("SELECT subject, body FROM templates WHERE name=?", (name,))
    if not row:
        return {"error": f"no template '{name}'"}
    subject = Template(row["subject"] or "").render(**context)
    body = Template(row["body"]).render(**context)
    return {"subject": subject, "body": body}


if __name__ == "__main__":
    mcp.run()
